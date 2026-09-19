import io
import csv
from typing import Optional
from fastapi import APIRouter, Depends, Header, Query, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.database import get_db
from app.models.school import ClassModel
from app.models.submission import SubmissionModel
from app.models.test import AssembledTestModel

router = APIRouter(prefix="/reports", tags=["Gradebook Exports"])

@router.get(
    "/assignments/{assignment_id}/gradebook.xlsx",
    summary="Export Electronic Gradebook",
    description="Downloads an Excel (.xlsx) or CSV spreadsheet ready for direct upload into school electronic gradebooks (e.g., edu.tatar.ru or МЭШ).",
)
async def export_gradebook(
    assignment_id: str,
    class_id: str = Query(..., description="Class ID, e.g. cls_7a_2026"),
    format: str = Query("xlsx", description="Output format: xlsx or csv", pattern="^(xlsx|csv)$"),
    x_teacher_uuid: Optional[str] = Header(None, alias="X-Teacher-UUID"),
    db: AsyncSession = Depends(get_db),
):
    stmt_test = select(AssembledTestModel).where(AssembledTestModel.test_id == assignment_id)
    res_test = await db.execute(stmt_test)
    test_obj = res_test.scalar_one_or_none()
    test_title = test_obj.title if test_obj else assignment_id

    stmt_cls = select(ClassModel).where(ClassModel.class_id == class_id)
    res_cls = await db.execute(stmt_cls)
    cls_obj = res_cls.scalar_one_or_none()
    class_name = cls_obj.name if cls_obj else class_id

    stmt_sub = (
        select(SubmissionModel)
        .where(SubmissionModel.assignment_id == assignment_id, SubmissionModel.class_id == class_id)
        .order_by(SubmissionModel.student_name)
    )
    res_sub = await db.execute(stmt_sub)
    submissions = res_sub.scalars().all()

    headers = [
        "ФИО ученика",
        "Вариант",
        "Набранный балл",
        "Макс. балл",
        "Оценка",
        "Дата проверки",
    ]

    rows = []
    for s in submissions:
        date_str = s.checked_at.strftime("%d.%m.%Y %H:%M") if s.checked_at else ""
        rows.append([
            s.student_name,
            s.variant,
            round(s.overall_score, 1),
            round(s.max_score, 1),
            s.final_grade,
            date_str,
        ])

    if format == "csv":
        output = io.StringIO()
        # UTF-8 BOM for Excel compatibility with Tatar Cyrillic
        output.write("\ufeff")
        writer = csv.writer(output, delimiter=";")
        writer.writerow([f"Контрольная работа: {test_title} | Класс: {class_name}"])
        writer.writerow([])
        writer.writerow(headers)
        for r in rows:
            writer.writerow(r)

        output.seek(0)
        csv_bytes = output.getvalue().encode("utf-8")
        filename = f"gradebook_{class_id}_{assignment_id}.csv"
        return StreamingResponse(
            io.BytesIO(csv_bytes),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    # Excel (.xlsx) output
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Журнал {class_name}"

    # Document Header Title
    title_font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
    ws.merge_cells("A1:F1")
    ws["A1"] = f"«Дәресханә» — Ведомость оценок: {test_title} ({class_name})"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A2"] = f"Сыйныф: {class_name} | Эш коды: {assignment_id} | Барлыгы укучылар: {len(submissions)}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="555555")

    # Table Header Row
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 24
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    zebra_fill = PatternFill(start_color="F9FBF9", end_color="F9FBF9", fill_type="solid")
    grade_5_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    grade_2_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    for r_idx, row_data in enumerate(rows, start=header_row_idx + 1):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        grade_val = row_data[4]

        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=11)

            # Center alignment for numbers and date
            if c_idx in (2, 3, 4, 5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Highlighting
            if c_idx == 5:
                if grade_val == 5:
                    cell.fill = grade_5_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="1B5E20")
                elif grade_val == 2:
                    cell.fill = grade_2_fill
                    cell.font = Font(name="Calibri", size=11, bold=True, color="B71C1C")
                else:
                    cell.fill = zebra_fill if is_even else PatternFill(fill_type=None)
            elif is_even:
                cell.fill = zebra_fill

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row > 2:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)

    filename = f"gradebook_{class_id}_{assignment_id}.xlsx"
    return StreamingResponse(
        excel_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
